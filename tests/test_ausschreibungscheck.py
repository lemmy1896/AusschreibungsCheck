import datetime as dt
import io
import json
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import ausschreibungscheck as ac  # noqa: E402

PROFIL = Path(__file__).resolve().parents[1] / "profil_ifs.yaml"


def release(ocid, titel, beschreibung="", cpv="71322000", region="DE929", plattform="evergabe-online.de"):
    return {
        "ocid": ocid,
        "date": "2026-10-12T08:00:00Z",
        "parties": [{
            "id": "b1", "name": "Stadt Hannover", "roles": ["buyer"],
            "address": {"postalCode": "30159", "locality": "Hannover", "region": region},
        }],
        "tender": {
            "title": titel,
            "description": beschreibung,
            "items": [{"classification": {"scheme": "CPV", "id": cpv}}],
            "tenderPeriod": {"endDate": "2026-11-05T12:00:00+01:00"},
            "submissionMethodDetails": f"https://www.{plattform}/tenderdetails.html?id=123",
        },
    }


def export_zip(releases):
    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w") as zf:
        zf.writestr("export.json", json.dumps({"releases": releases}))
    return puffer.getvalue()


def test_bewertung_trifft_fachthema():
    profil = ac.Profil.laden(PROFIL)
    b = ac.aus_release(release("1", "Generalentwässerungsplanung Stadt Laatzen", "hydraulische Berechnung des Kanalnetzes"))
    t = ac.bewerten(b, profil)
    assert t is not None
    assert "Generalentwässerungsplanung" in t.gruende
    assert "CPV 71322000" in t.gruende
    assert "Region DE9" in t.gruende
    assert b.link.startswith("https://www.evergabe-online.de/")
    assert b.ort == "30159 Hannover"
    assert b.frist == "2026-11-05 12:00"


def test_fremdes_thema_und_ausschluss():
    profil = ac.Profil.laden(PROFIL)
    assert ac.bewerten(ac.aus_release(release("2", "Lieferung von Büromöbeln", cpv="39130000")), profil) is None
    assert ac.bewerten(ac.aus_release(release("3", "Kanalreinigung Regenwasserkanalnetz")), profil) is None


def test_abkuerzung_nur_als_wort():
    profil = ac.Profil.laden(PROFIL)
    b = ac.aus_release(release("4", "GEPÄCKBÄNDER Flughafen", cpv="42000000", region="DE1"))
    assert ac.bewerten(b, profil) is None
    b = ac.aus_release(release("5", "Fortschreibung GEP und Fremdwasserkonzept", cpv="42000000", region="DE1"))
    assert ac.bewerten(b, profil) is not None


def test_plattformfilter():
    profil = ac.Profil.laden(PROFIL)
    assert ac.passt_plattform(ac.aus_release(release("6", "x")), profil)
    assert not ac.passt_plattform(ac.aus_release(release("7", "x", plattform="vergabe.niedersachsen.de")), profil)


def test_tage_nachholen():
    heute = dt.date(2026, 10, 13)
    assert len(ac.zu_pruefende_tage({"letzter_tag": None}, heute, None)) == 30
    assert ac.zu_pruefende_tage({"letzter_tag": "2026-10-10"}, heute, None) == [dt.date(2026, 10, 11), dt.date(2026, 10, 12)]
    assert ac.zu_pruefende_tage({"letzter_tag": "2026-10-12"}, heute, None) == []


def test_gesamtlauf(tmp_path, monkeypatch):
    (tmp_path / "profil_ifs.yaml").write_text(PROFIL.read_text(encoding="utf-8"), encoding="utf-8")
    zipdatei = tmp_path / "2026-10-12.zip"
    zipdatei.write_bytes(export_zip([
        release("a", "Retentionsbodenfilter Volksparkstraße, Objektplanung LPH 1-6"),
        release("b", "Abflussmessungen Dattelner Mühlenbach", cpv="71351000", region="DEA36"),
        release("c", "Regenrückhaltebecken Nord", plattform="vergabe.niedersachsen.de"),
        release("d", "Reinigung Verwaltungsgebäude", cpv="90910000"),
    ]))
    monkeypatch.setattr(ac, "basisordner", lambda: tmp_path)
    assert ac.main(["--zip", str(zipdatei), "--nicht-oeffnen"]) == 0
    listen = list((tmp_path / "ergebnisse").glob("*.xlsx"))
    assert len(listen) == 1
    ws = load_workbook(listen[0])["Treffer"]
    zeilen = list(ws.iter_rows(min_row=2, values_only=True))
    assert {z[2] for z in zeilen} == {"Retentionsbodenfilter Volksparkstraße, Objektplanung LPH 1-6",
                                      "Abflussmessungen Dattelner Mühlenbach"}
    assert all(z[0] == "neu" for z in zeilen)

    # zweiter Lauf: dieselben Treffer sind nicht mehr neu
    for f in listen:
        f.unlink()
    assert ac.main(["--zip", str(zipdatei), "--nicht-oeffnen"]) == 0
    ws = load_workbook(next((tmp_path / "ergebnisse").glob("*.xlsx")))["Treffer"]
    assert all(z[0] is None for z in ws.iter_rows(min_row=2, values_only=True))
