"""AusschreibungsCheck: sucht passende öffentliche Ausschreibungen.

Lädt die offenen Tagesexporte (OCDS) des Bekanntmachungsservice des Bundes
(oeffentlichevergabe.de), bewertet jede Bekanntmachung anhand eines
YAML-Suchprofils und schreibt die Treffer in eine Excel-Liste.
"""

from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
import re
import sys
import urllib.error
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_URL = "https://oeffentlichevergabe.de/api/notice-exports?pubDay={day}&format={format}"
NOTICE_URL = "https://oeffentlichevergabe.de/ui/de/search?query={query}"
ERSTLAUF_TAGE = 30
USER_AGENT = "AusschreibungsCheck/1.0"


def basisordner() -> Path:
    """Ordner der .exe bzw. des Skripts; dort liegen Profil, Daten und Ergebnisse."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


# --------------------------------------------------------------------------- Profil

@dataclass
class Profil:
    name: str
    plattformen: dict[str, list[str]]  # Anzeigename -> Hosts der Plattform
    mindestpunkte: int
    begriffe: dict[str, int]
    cpv: dict[str, int]
    regionen: dict[str, int]
    ausschluss: list[str]
    _muster: list[tuple[str, re.Pattern, int]] = field(default_factory=list, repr=False)

    @classmethod
    def laden(cls, pfad: Path) -> "Profil":
        daten = yaml.safe_load(pfad.read_text(encoding="utf-8")) or {}
        profil = cls(
            name=str(daten.get("name", pfad.stem)),
            plattformen=plattformen_lesen(daten.get("plattformen")),
            mindestpunkte=int(daten.get("mindestpunkte", 10)),
            begriffe={str(k): int(v) for k, v in (daten.get("begriffe") or {}).items()},
            cpv={str(k): int(v) for k, v in (daten.get("cpv") or {}).items()},
            regionen={str(k).upper(): int(v) for k, v in (daten.get("regionen") or {}).items()},
            ausschluss=[str(a) for a in daten.get("ausschluss") or []],
        )
        profil._muster_bauen()
        return profil

    def _muster_bauen(self) -> None:
        self._muster = [(b, begriff_muster(b), p) for b, p in self.begriffe.items()]

    def mit_zusatzbegriffen(self, zusatz: dict[str, int]) -> "Profil":
        """Ergänzt Schlagwörter aus dem Programmfenster; vorhandene Gewichte werden überschrieben."""
        self.begriffe = {**self.begriffe, **zusatz}
        self._muster_bauen()
        return self


def plattformen_lesen(eintrag) -> dict[str, list[str]]:
    if not eintrag:
        return {}
    if isinstance(eintrag, list):
        return {str(h): [str(h).lower()] for h in eintrag}
    return {str(name): [str(h).lower() for h in (hosts if isinstance(hosts, list) else [hosts])]
            for name, hosts in eintrag.items()}


def zusatzbegriffe_lesen(text: str, standardgewicht: int) -> dict[str, int]:
    """Eine Zeile pro Schlagwort, optional mit Gewicht: "Starkregen" oder "Starkregen: 12"."""
    begriffe: dict[str, int] = {}
    for zeile in text.splitlines():
        zeile = zeile.split("#", 1)[0].strip()
        if not zeile:
            continue
        treffer = re.match(r"^(.*?)\s*[:=]\s*(\d+)$", zeile)
        if treffer and treffer.group(1):
            begriffe[treffer.group(1)] = int(treffer.group(2))
        else:
            begriffe[zeile] = standardgewicht
    return begriffe


def begriff_muster(begriff: str) -> re.Pattern:
    """Abkürzungen in Großbuchstaben (GEP, WRRL) nur als ganzes Wort, sonst auch in Komposita."""
    text = re.escape(begriff)
    if begriff.isupper():
        return re.compile(rf"(?<![\wÄÖÜäöüß]){text}(?![\wÄÖÜäöüß])")
    return re.compile(text, re.IGNORECASE)


# --------------------------------------------------------------------------- Bekanntmachungen

@dataclass
class Bekanntmachung:
    id: str
    titel: str
    beschreibung: str
    vergabestelle: str
    ort: str
    regionen: list[str]
    cpv: list[str]
    frist: str
    veroeffentlicht: str
    links: list[str]
    dokument_links: list[str]
    rohtext: str

    plattform: str = ""

    @property
    def link(self) -> str:
        if self.dokument_links:
            return self.dokument_links[0]
        return NOTICE_URL.format(query=urllib.request.quote(self.titel[:80]))

    @property
    def hosts(self) -> set[str]:
        return {re.sub(r"^https?://([^/:?#]+).*", r"\1", u.strip()).lower() for u in self.dokument_links}


def alle_dicts(obj):
    if isinstance(obj, dict):
        yield obj
        for wert in obj.values():
            yield from alle_dicts(wert)
    elif isinstance(obj, list):
        for wert in obj:
            yield from alle_dicts(wert)


def alle_urls(obj) -> list[str]:
    urls: list[str] = []
    for d in alle_dicts(obj):
        for wert in d.values():
            if isinstance(wert, str) and wert.startswith(("http://", "https://")) and wert not in urls:
                urls.append(wert)
    return urls


def aus_release(release: dict) -> Bekanntmachung:
    tender = release.get("tender") or {}
    parties = release.get("parties") or []

    kaeufer = next((p for p in parties if "buyer" in (p.get("roles") or [])), None)
    if kaeufer is None:
        kaeufer = (release.get("buyer") or {}) | {}
        kaeufer = next((p for p in parties if p.get("id") == kaeufer.get("id")), kaeufer)
    adresse = (kaeufer or {}).get("address") or {}

    cpv: list[str] = []
    for item in tender.get("items") or []:
        for klass in [item.get("classification") or {}] + (item.get("additionalClassifications") or []):
            if str(klass.get("scheme", "CPV")).upper() == "CPV" and klass.get("id"):
                code = str(klass["id"]).split("-")[0]
                if code not in cpv:
                    cpv.append(code)
    if (tender.get("classification") or {}).get("id"):
        code = str(tender["classification"]["id"]).split("-")[0]
        if code not in cpv:
            cpv.insert(0, code)

    regionen: list[str] = []
    for d in alle_dicts(release):
        for schluessel in ("region", "nutsCode", "NUTS"):
            wert = d.get(schluessel)
            if isinstance(wert, str) and re.match(r"^DE[0-9A-G]", wert.upper()) and wert.upper() not in regionen:
                regionen.append(wert.upper())

    frist = ((tender.get("tenderPeriod") or {}).get("endDate")) or release.get("_frist") or ""
    dokument_links = [d["url"] for d in tender.get("documents") or [] if isinstance(d, dict) and d.get("url")]
    if str(tender.get("submissionMethodDetails", "")).startswith("http"):
        dokument_links.append(tender["submissionMethodDetails"])
    return Bekanntmachung(
        id=str(release.get("ocid") or release.get("id") or ""),
        titel=str(tender.get("title") or release.get("title") or "").strip(),
        beschreibung=str(tender.get("description") or "").strip(),
        vergabestelle=str((kaeufer or {}).get("name") or ""),
        ort=" ".join(x for x in (adresse.get("postalCode"), adresse.get("locality")) if x),
        regionen=regionen,
        cpv=cpv,
        frist=frist[:16].replace("T", " "),
        veroeffentlicht=str(release.get("date") or "")[:10],
        links=alle_urls(release),
        dokument_links=dokument_links,
        rohtext=json.dumps(release, ensure_ascii=False).lower(),
    )


def releases_aus_json(daten) -> list[dict]:
    if isinstance(daten, list):
        return [r for eintrag in daten for r in releases_aus_json(eintrag)]
    if not isinstance(daten, dict):
        return []
    if "releases" in daten:
        return [r for r in daten["releases"] if isinstance(r, dict)]
    if "records" in daten:
        return [(r.get("compiledRelease") or r) for r in daten["records"] if isinstance(r, dict)]
    if "tender" in daten or "ocid" in daten:
        return [daten]
    return []


def releases_aus_zip(inhalt: bytes, fristen: dict[str, str] | None = None) -> list[dict]:
    """Liest alle Releases eines OCDS-Exports; Fristen aus dem eForms-Export werden ergänzt."""
    releases: list[dict] = []
    with zipfile.ZipFile(io.BytesIO(inhalt)) as zf:
        for name in zf.namelist():
            if not name.lower().endswith(".json"):
                continue
            with zf.open(name) as f:
                try:
                    neue = releases_aus_json(json.load(f))
                except json.JSONDecodeError:
                    print(f"  Warnung: {name} ist kein gültiges JSON und wird übersprungen.")
                    continue
            frist = (fristen or {}).get(Path(name).stem)
            for r in neue:
                if frist:
                    r["_frist"] = frist
            releases.extend(neue)
    return releases


def lokaler_name(element: ET.Element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def frist_aus_eforms(xml: bytes) -> str | None:
    """Früheste Angebotsfrist (TenderSubmissionDeadlinePeriod) einer eForms-Bekanntmachung."""
    try:
        wurzel = ET.fromstring(xml)
    except ET.ParseError:
        return None
    fristen = []
    for periode in wurzel.iter():
        if lokaler_name(periode) != "TenderSubmissionDeadlinePeriod":
            continue
        werte = {lokaler_name(k): (k.text or "").strip() for k in periode}
        datum = werte.get("EndDate", "")[:10]
        if datum:
            fristen.append(f"{datum}T{werte.get('EndTime', '')[:5]}".rstrip("T"))
    return min(fristen) if fristen else None


def fristen_aus_zip(inhalt: bytes) -> dict[str, str]:
    fristen: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(inhalt)) as zf:
        for name in zf.namelist():
            if name.lower().endswith(".xml") and (frist := frist_aus_eforms(zf.read(name))):
                fristen[Path(name).stem] = frist
    return fristen


# --------------------------------------------------------------------------- Abruf

class AbrufFehler(Exception):
    """Netzwerkproblem: dieser und spätere Tage werden beim nächsten Lauf erneut versucht."""


def tagesexport(tag: dt.date, cache: Path, format_: str = "ocds.zip", log=print) -> bytes | None:
    datei = cache / f"{tag.isoformat()}.{format_}"
    if datei.exists() and datei.stat().st_size > 0:
        return datei.read_bytes()
    url = EXPORT_URL.format(day=tag.isoformat(), format=format_)
    anfrage = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(anfrage, timeout=120) as antwort:
            inhalt = antwort.read()
    except urllib.error.HTTPError as fehler:
        log(f"  {tag}: kein Export verfügbar (HTTP {fehler.code}).")
        return None
    except (urllib.error.URLError, TimeoutError) as fehler:
        raise AbrufFehler(f"{tag}: Abruf fehlgeschlagen ({fehler})") from fehler
    cache.mkdir(parents=True, exist_ok=True)
    datei.write_bytes(inhalt)
    return inhalt


# --------------------------------------------------------------------------- Bewertung

@dataclass
class Treffer:
    bekanntmachung: Bekanntmachung
    punkte: int
    gruende: list[str]
    neu: bool = True


def plattform_von(b: Bekanntmachung, plattformen: dict[str, list[str]]) -> str | None:
    """Die Plattform steht im Link zu den Vergabeunterlagen (tender.documents)."""
    for name, hosts in plattformen.items():
        for host in hosts:
            if any(h == host or h.endswith("." + host) for h in b.hosts):
                return name
    return None


def passt_plattform(b: Bekanntmachung, profil: Profil) -> bool:
    return not profil.plattformen or plattform_von(b, profil.plattformen) is not None


def bewerten(b: Bekanntmachung, profil: Profil) -> Treffer | None:
    text = f"{b.titel}\n{b.beschreibung}"
    text_klein = text.lower()
    if any(a.lower() in text_klein for a in profil.ausschluss):
        return None

    punkte = 0
    gruende: list[str] = []
    for begriff, muster, wert in profil._muster:
        if muster.search(text):
            punkte += wert
            gruende.append(begriff)
    for code, wert in profil.cpv.items():
        praefix = code.rstrip("0") or code
        if any(c.startswith(praefix) for c in b.cpv):
            punkte += wert
            gruende.append(f"CPV {code}")
    for region, wert in profil.regionen.items():
        if any(r.startswith(region) for r in b.regionen):
            punkte += wert
            gruende.append(f"Region {region}")
            break

    if punkte < profil.mindestpunkte:
        return None
    return Treffer(b, punkte, gruende)


# --------------------------------------------------------------------------- Zustand

def zustand_laden(pfad: Path) -> dict:
    if pfad.exists():
        return json.loads(pfad.read_text(encoding="utf-8"))
    return {"letzter_tag": None, "gesehen": []}


def zustand_speichern(pfad: Path, zustand: dict) -> None:
    pfad.write_text(json.dumps(zustand, ensure_ascii=False, indent=1), encoding="utf-8")


def zu_pruefende_tage(zustand: dict, heute: dt.date, von: dt.date | None) -> list[dt.date]:
    bis = heute - dt.timedelta(days=1)  # der heutige Export ist noch unvollständig
    if von is None:
        if zustand.get("letzter_tag"):
            von = dt.date.fromisoformat(zustand["letzter_tag"]) + dt.timedelta(days=1)
        else:
            von = heute - dt.timedelta(days=ERSTLAUF_TAGE)
    tage = []
    tag = von
    while tag <= bis:
        tage.append(tag)
        tag += dt.timedelta(days=1)
    return tage


# --------------------------------------------------------------------------- Excel

SPALTEN = [
    ("Neu", 6), ("Punkte", 8), ("Titel", 60), ("Plattform", 18), ("Vergabestelle", 35), ("Ort", 22),
    ("Frist", 17), ("Veröffentlicht", 13), ("CPV", 22), ("Treffergründe", 45), ("Link", 50),
]


def excel_schreiben(treffer: list[Treffer], pfad: Path, profilname: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Treffer"
    ws.append([s for s, _ in SPALTEN])
    for zelle in ws[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill("solid", fgColor="1F4E78")
    neu_fill = PatternFill("solid", fgColor="E2EFDA")

    for t in treffer:
        b = t.bekanntmachung
        ws.append([
            "neu" if t.neu else "", t.punkte, b.titel, b.plattform, b.vergabestelle, b.ort, b.frist,
            b.veroeffentlicht, ", ".join(b.cpv), ", ".join(t.gruende), b.link,
        ])
        zeile = ws.max_row
        link = ws.cell(row=zeile, column=len(SPALTEN))
        link.hyperlink = b.link
        link.font = Font(color="0563C1", underline="single")
        if t.neu:
            for spalte in range(1, len(SPALTEN) + 1):
                ws.cell(row=zeile, column=spalte).fill = neu_fill

    for i, (_, breite) in enumerate(SPALTEN, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Info")
    info.append(["Profil", profilname])
    info.append(["Erstellt", dt.datetime.now().strftime("%d.%m.%Y %H:%M")])
    info.append(["Treffer", len(treffer)])
    info.append(["Neue Treffer", sum(t.neu for t in treffer)])
    wb.save(pfad)


# --------------------------------------------------------------------------- Ablauf

def profil_finden(basis: Path) -> Path | None:
    profile = sorted(basis.glob("profil_*.yaml"))
    return profile[0] if profile else None


def ausfuehren(
    profilpfad: Path,
    basis: Path,
    plattformen: list[str] | None = None,
    alle_plattformen: bool = False,
    zusatzbegriffe: dict[str, int] | None = None,
    von: dt.date | None = None,
    zips: list[Path] | None = None,
    log=print,
) -> Path:
    """Führt eine Suche aus und gibt den Pfad der Excel-Liste zurück.

    plattformen: Anzeigenamen aus dem Profil (None = alle im Profil genannten).
    alle_plattformen: auch Bekanntmachungen anderer oder unbekannter Plattformen.
    """
    profil = Profil.laden(profilpfad)
    if zusatzbegriffe:
        profil.mit_zusatzbegriffen(zusatzbegriffe)
    auswahl = {n: h for n, h in profil.plattformen.items() if plattformen is None or n in plattformen}
    log(f"Profil: {profil.name} ({profilpfad.name})")

    datenordner = basis / "daten"
    datenordner.mkdir(exist_ok=True)
    zustandspfad = datenordner / f"zustand_{profilpfad.stem}.json"
    zustand = zustand_laden(zustandspfad)
    gesehen = set(zustand.get("gesehen", []))

    releases: list[dict] = []
    tage: list[dt.date] = []
    if zips:
        for z in zips:
            log(f"Lese {z} ...")
            releases.extend(releases_aus_zip(Path(z).read_bytes()))
    else:
        tage = zu_pruefende_tage(zustand, dt.date.today(), von)
        if not tage:
            log("Seit dem letzten Lauf ist kein neuer Tag dazugekommen.")
        for i, tag in enumerate(tage):
            log(f"Lade Bekanntmachungen vom {tag:%d.%m.%Y} ...")
            try:
                inhalt = tagesexport(tag, datenordner / "cache", log=log)
            except AbrufFehler as fehler:
                log(f"  {fehler}. Ab hier wird beim nächsten Lauf erneut geladen.")
                tage = tage[:i]
                break
            if inhalt:
                fristen: dict[str, str] = {}
                try:
                    eforms = tagesexport(tag, datenordner / "cache", "eforms.zip", log=log)
                    fristen = fristen_aus_zip(eforms) if eforms else {}
                except (AbrufFehler, zipfile.BadZipFile) as fehler:
                    log(f"  {tag}: Fristen nicht verfügbar ({fehler}).")
                try:
                    releases.extend(releases_aus_zip(inhalt, fristen))
                except zipfile.BadZipFile:
                    log(f"  {tag}: Export ist beschädigt und wird übersprungen.")
                    (datenordner / "cache" / f"{tag.isoformat()}.ocds.zip").unlink(missing_ok=True)

    bekanntmachungen: dict[str, Bekanntmachung] = {}
    for r in releases:
        b = aus_release(r)
        bekanntmachungen[b.id or b.titel] = b
    ausgewaehlt = []
    for b in bekanntmachungen.values():
        b.plattform = plattform_von(b, profil.plattformen) or ""
        if alle_plattformen or not profil.plattformen or plattform_von(b, auswahl):
            ausgewaehlt.append(b)
    beschreibung = "allen Plattformen" if alle_plattformen else ", ".join(auswahl) or "keiner Plattform"
    log(f"{len(bekanntmachungen)} Bekanntmachungen gelesen, davon {len(ausgewaehlt)} von {beschreibung}.")

    heute = dt.date.today().isoformat()
    offen = [b for b in ausgewaehlt if not b.frist or b.frist[:10] >= heute]
    if len(offen) < len(ausgewaehlt):
        log(f"{len(ausgewaehlt) - len(offen)} Bekanntmachungen mit abgelaufener Frist ausgelassen.")
    treffer = [t for b in offen if (t := bewerten(b, profil))]
    for t in treffer:
        t.neu = t.bekanntmachung.id not in gesehen
    treffer.sort(key=lambda t: (not t.neu, -t.punkte))

    ergebnisse = basis / "ergebnisse"
    ergebnisse.mkdir(exist_ok=True)
    excel = ergebnisse / f"Treffer_{profilpfad.stem}_{dt.datetime.now():%Y-%m-%d_%H%M%S}.xlsx"
    excel_schreiben(treffer, excel, profil.name)
    log(f"{len(treffer)} Treffer, davon {sum(t.neu for t in treffer)} neu.")
    log(f"Liste: {excel}")

    gesehen.update(t.bekanntmachung.id for t in treffer)
    zustand["gesehen"] = sorted(gesehen)
    if tage:
        zustand["letzter_tag"] = max(tage[-1].isoformat(), zustand.get("letzter_tag") or "")
    zustand_speichern(zustandspfad, zustand)
    return excel


def oeffnen(pfad: Path) -> None:
    if hasattr(os, "startfile"):
        os.startfile(pfad)  # type: ignore[attr-defined]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Sucht passende Ausschreibungen.")
    parser.add_argument("--profil", help="Pfad zur Profil-Datei (Standard: profil_*.yaml im Programmordner)")
    parser.add_argument("--von", help="Ab diesem Tag prüfen (JJJJ-MM-TT), statt seit dem letzten Lauf")
    parser.add_argument("--plattform", action="append", help="Nur diese Plattform (Name aus dem Profil), mehrfach möglich")
    parser.add_argument("--alle-plattformen", action="store_true", help="Auch andere und unbekannte Plattformen")
    parser.add_argument("--begriff", action="append", default=[], help='Zusätzliches Schlagwort, z. B. "Starkregen: 12"')
    parser.add_argument("--zip", nargs="*", help="Lokale Export-Dateien statt Download verwenden (zum Testen)")
    parser.add_argument("--nicht-oeffnen", action="store_true", help="Excel-Liste am Ende nicht öffnen")
    args = parser.parse_args(argv)

    basis = basisordner()
    profilpfad = Path(args.profil) if args.profil else profil_finden(basis)
    if profilpfad is None:
        print(f"Keine Profil-Datei (profil_*.yaml) in {basis} gefunden.")
        return 1
    mindest = Profil.laden(profilpfad).mindestpunkte
    excel = ausfuehren(
        profilpfad, basis,
        plattformen=args.plattform,
        alle_plattformen=args.alle_plattformen,
        zusatzbegriffe=zusatzbegriffe_lesen("\n".join(args.begriff), mindest),
        von=dt.date.fromisoformat(args.von) if args.von else None,
        zips=[Path(z) for z in args.zip] if args.zip else None,
    )
    if not args.nicht_oeffnen:
        oeffnen(excel)
    return 0


if __name__ == "__main__":
    sys.exit(main())
